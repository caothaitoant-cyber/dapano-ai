# -*- coding: utf-8 -*-
"""Dựng Máy Tính Dòng Tiền 10 Năm - quà Tầng 1 DAPANO."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter

INDIGO="FF21135F"; PURPLE="FF432A73"; MAGENTA="FF5E067C"; GOLD="FFD4A856"
GOLDL="FFF0E0BA"; BGSOFT="FFF5F1F8"; CREAM="FFFAF7FC"; LIGHT="FFE8E4F0"
TDARK="FF1A0F4A"; TSOFT="FF4A3D6E"; TMUTE="FF6B5E85"; WHITE="FFFFFFFF"

IN = "'1. Nhập liệu'"
M0 = '#,##0'; PCT = '0.0%'; NUM2 = '0.00'

thin = Side(style='thin', color=LIGHT)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
gold_box = Border(left=Side(style='thin',color=GOLD), right=Side(style='thin',color=GOLD),
                  top=Side(style='thin',color=GOLD), bottom=Side(style='thin',color=GOLD))

wb = Workbook()

def title_block(ws, title, sub, span=6):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=span)
    c = ws.cell(row=1,column=1,value=title)
    c.font = Font(name='Montserrat',size=16,bold=True,color=WHITE)
    c.fill = PatternFill('solid', fgColor=INDIGO)
    c.alignment = Alignment(horizontal='left',vertical='center',indent=1)
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=span)
    c2 = ws.cell(row=2,column=1,value=sub)
    c2.font = Font(name='Montserrat',size=9,italic=True,color=WHITE)
    c2.fill = PatternFill('solid', fgColor=MAGENTA)
    c2.alignment = Alignment(horizontal='left',vertical='center',indent=1)
    ws.row_dimensions[2].height = 20
    for col in range(1,span+1):
        ws.cell(row=1,column=col).fill = PatternFill('solid', fgColor=INDIGO)
        ws.cell(row=2,column=col).fill = PatternFill('solid', fgColor=MAGENTA)

def section(ws, row, text, span=4):
    ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=span)
    c = ws.cell(row=row,column=1,value=text)
    c.font = Font(name='Montserrat',size=10,bold=True,color=WHITE)
    c.alignment = Alignment(horizontal='left',vertical='center',indent=1)
    for col in range(1,span+1):
        ws.cell(row=row,column=col).fill = PatternFill('solid', fgColor=PURPLE)
    ws.row_dimensions[row].height = 22

def label(ws,row,text,note=""):
    c = ws.cell(row=row,column=1,value=text)
    c.font = Font(name='Montserrat',size=10,color=TDARK)
    c.alignment = Alignment(vertical='center',indent=1,wrap_text=True)
    n = ws.cell(row=row,column=4,value=note)
    n.font = Font(name='Montserrat',size=8,italic=True,color=TMUTE)
    n.alignment = Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[row].height = 20

def input_cell(ws,row,unit,fmt=M0,default=None):
    c = ws.cell(row=row,column=2,value=default)
    c.fill = PatternFill('solid', fgColor=GOLDL)
    c.font = Font(name='Montserrat',size=11,bold=True,color=TDARK)
    c.number_format = fmt
    c.border = gold_box
    c.alignment = Alignment(horizontal='right',vertical='center')
    c.protection = Protection(locked=False)
    u = ws.cell(row=row,column=3,value=unit)
    u.font = Font(name='Montserrat',size=9,color=TMUTE)
    u.alignment = Alignment(vertical='center',indent=1)
    return c

def calc_cell(ws,row,formula,unit,fmt=M0,col=2,bold=True):
    c = ws.cell(row=row,column=col,value=formula)
    c.fill = PatternFill('solid', fgColor=BGSOFT)
    c.font = Font(name='Montserrat',size=11,bold=bold,color=MAGENTA)
    c.number_format = fmt
    c.border = box
    c.alignment = Alignment(horizontal='right',vertical='center')
    if unit is not None:
        u = ws.cell(row=row,column=3,value=unit)
        u.font = Font(name='Montserrat',size=9,color=TMUTE)
        u.alignment = Alignment(vertical='center',indent=1)
    return c

# ═══════════════════ SHEET 1 · NHẬP LIỆU ═══════════════════
ws = wb.active; ws.title = "1. Nhập liệu"
title_block(ws,"MÁY TÍNH DÒNG TIỀN 10 NĂM","DAPANO GROUP · Cash Flow First · Chỉ gõ vào những ô màu vàng",4)
ws.column_dimensions['A'].width = 42
ws.column_dimensions['B'].width = 16
ws.column_dimensions['C'].width = 14
ws.column_dimensions['D'].width = 46

ws.merge_cells('A4:D4')
g = ws.cell(row=4,column=1,value="Đơn vị tiền tệ dùng trong toàn bộ file: TRIỆU ĐỒNG. Ví dụ toà nhà 12 tỷ thì gõ 12000.")
g.font = Font(name='Montserrat',size=9,bold=True,color=MAGENTA)
g.fill = PatternFill('solid', fgColor=CREAM); g.alignment = Alignment(indent=1,vertical='center')
ws.row_dimensions[4].height = 20

section(ws,6,"A · TIỀN BỎ RA ĐỂ SỞ HỮU")
label(ws,7,"Giá mua toà nhà","Số ghi trên hợp đồng công chứng")
input_cell(ws,7,"triệu đồng",default=12000)
label(ws,8,"Thuế, phí, công chứng","Thường 2–3% giá mua")
input_cell(ws,8,"triệu đồng",default=300)
label(ws,9,"Phí môi giới","Bỏ trống nếu mua trực tiếp")
input_cell(ws,9,"triệu đồng",default=120)
label(ws,10,"Chi phí cải tạo trước khi cho thuê","Điện nước, chống thấm, sơn, thang máy")
input_cell(ws,10,"triệu đồng",default=800)
label(ws,11,"TỔNG VỐN ĐẦU TƯ","Đây mới là số tiền thật anh/chị bỏ vào tài sản")
calc_cell(ws,11,"=SUM(B7:B10)","triệu đồng")

section(ws,13,"B · CẤU TRÚC VỐN")
label(ws,14,"Vốn tự có","Tiền của mình, không phải tiền vay")
input_cell(ws,14,"triệu đồng",default=9500)
label(ws,15,"Số tiền vay ngân hàng","Tự tính = tổng vốn − vốn tự có")
calc_cell(ws,15,"=MAX(0,B11-B14)","triệu đồng")
label(ws,16,"Tỷ lệ vay trên tổng vốn","Trên 70% là vùng rủi ro cao")
calc_cell(ws,16,"=IF(B11<=0,0,B15/B11)","",fmt=PCT)
label(ws,17,"Lãi suất năm đầu (ưu đãi)","Gõ 8,5% thì nhập 8.5%")
input_cell(ws,17,"%/năm",fmt=PCT,default=0.085)
label(ws,18,"Lãi suất thả nổi các năm sau","Hỏi ngân hàng biên độ cộng thêm là bao nhiêu")
input_cell(ws,18,"%/năm",fmt=PCT,default=0.115)
label(ws,19,"Kỳ hạn vay","Trả gốc đều, lãi tính trên dư nợ giảm dần")
input_cell(ws,19,"năm",fmt='0',default=20)

section(ws,21,"C · VẬN HÀNH CHO THUÊ")
label(ws,22,"Giá thuê thu được mỗi tháng","Giá đã ký được, không phải giá kỳ vọng")
input_cell(ws,22,"triệu/tháng",default=75)
label(ws,23,"Số tháng để trống mỗi năm","Thực tế thường 0,5–2 tháng. Không có toà nào bằng 0 mãi")
input_cell(ws,23,"tháng",fmt='0.0',default=1)
label(ws,24,"Chi phí vận hành mỗi tháng","Quản lý, sửa chữa, bảo hiểm, phí. Đừng để trống ô này")
input_cell(ws,24,"triệu/tháng",default=6)
label(ws,25,"Thuế cho thuê","Cá nhân cho thuê thường 10% doanh thu")
input_cell(ws,25,"% doanh thu",fmt=PCT,default=0.1)

section(ws,27,"D · GIẢ ĐỊNH DÀI HẠN")
label(ws,28,"Tốc độ tăng giá thuê mỗi năm","Đừng lạc quan quá. 3–5% là hợp lý")
input_cell(ws,28,"%/năm",fmt=PCT,default=0.04)
label(ws,29,"Tốc độ tăng chi phí vận hành mỗi năm","Chi phí luôn tăng nhanh hơn ta tưởng")
input_cell(ws,29,"%/năm",fmt=PCT,default=0.05)

ws.merge_cells('A31:D33')
w = ws.cell(row=31,column=1,value=("Nhập xong, mở sang sheet 2 để xem dòng tiền 10 năm, sheet 3 để xem ba kịch bản.\n"
    "Nếu chưa biết một con số nào đó, cứ để nguyên số mẫu rồi sửa sau — máy tính không vỡ.\n"
    "Em luôn sẵn sàng ngồi cùng anh/chị đọc lại kết quả này.  ·  DAPANO GROUP"))
w.font = Font(name='Montserrat',size=9,color=TSOFT)
w.fill = PatternFill('solid', fgColor=BGSOFT)
w.alignment = Alignment(wrap_text=True,vertical='top',indent=1)
ws.freeze_panes = "A5"
ws.sheet_view.showGridLines = False
ws.protection.sheet = True
ws.protection.enable()

# ═══════════════════ SHEET 2 · KẾT QUẢ ═══════════════════
r = wb.create_sheet("2. Kết quả 10 năm")
title_block(r,"KẾT QUẢ 10 NĂM","Toàn bộ ô ở sheet này tự tính — không cần gõ gì",14)
HDR = ["Năm","Doanh thu thuê","Chi phí vận hành","Thuế cho thuê","NOI","Dư nợ đầu năm",
       "Trả gốc","Trả lãi","Tổng trả nợ","Dòng tiền ròng","DSCR","Dòng tiền luỹ kế",
       "Cap Rate / giá mua","Cap Rate / tổng vốn"]
widths = [7,15,15,14,13,14,12,12,13,15,9,16,15,16]
for i,w_ in enumerate(widths,1): r.column_dimensions[get_column_letter(i)].width = w_

# --- khối tóm tắt
section(r,4,"TÓM TẮT — ĐỌC BA DÒNG NÀY TRƯỚC",14)
summary = [
 ("NOI năm 1 (lãi ròng trước trả nợ)","=E15",M0,"triệu/năm"),
 ("Cap Rate trên tổng vốn (năm 1)","=N15",PCT,""),
 ("Dòng tiền ròng năm 1","=J15",M0,"triệu/năm"),
 ("Dòng tiền ròng bình quân tháng (năm 1)","=J15/12",M0,"triệu/tháng"),
 ("DSCR năm 1","=K15",NUM2,""),
 ("Đèn cảnh báo DSCR",'=IF(I15<=0,"Không vay — không có rủi ro trả nợ",IF(E15/I15<1.2,"ĐÈN ĐỎ · DSCR dưới 1,2 — đòn bẩy quá căng",IF(E15/I15<1.5,"ĐÈN VÀNG · an toàn vừa đủ, đừng vay thêm","ĐÈN XANH · dòng tiền gánh được nợ")))',None,""),
 ("Dòng tiền chuyển dương từ năm thứ",'=IFERROR(MATCH(TRUE,INDEX($J$15:$J$24>0,0),0),"chưa dương trong 10 năm")','0',""),
 ("Dòng tiền luỹ kế 10 năm so với vốn tự có",'=IF('+IN+'!$B$14<=0,"—",L24/'+IN+'!$B$14)',PCT,"chưa tính phần tăng giá tài sản"),
 ("Tổng dòng tiền ròng 10 năm","=SUM(J15:J24)",M0,"triệu"),
]
row = 5
for lbl,f,fmt,unit in summary:
    r.merge_cells(start_row=row,start_column=1,end_row=row,end_column=4)
    c = r.cell(row=row,column=1,value=lbl)
    c.font = Font(name='Montserrat',size=10,color=TDARK); c.alignment = Alignment(indent=1,vertical='center')
    v = r.cell(row=row,column=5,value=f)
    v.font = Font(name='Montserrat',size=11,bold=True,color=MAGENTA)
    if fmt: v.number_format = fmt
    v.fill = PatternFill('solid', fgColor=BGSOFT); v.border = box
    v.alignment = Alignment(horizontal='right',vertical='center')
    if lbl.startswith("Đèn"):
        r.merge_cells(start_row=row,start_column=5,end_row=row,end_column=9)
        v.alignment = Alignment(horizontal='left',vertical='center',indent=1)
        v.font = Font(name='Montserrat',size=10,bold=True,color=INDIGO)
    else:
        u = r.cell(row=row,column=6,value=unit)
        u.font = Font(name='Montserrat',size=9,color=TMUTE); u.alignment = Alignment(indent=1,vertical='center')
    r.row_dimensions[row].height = 19
    row += 1

# --- bảng 10 năm
HR = 14
for i,h in enumerate(HDR,1):
    c = r.cell(row=HR,column=i,value=h)
    c.font = Font(name='Montserrat',size=9,bold=True,color=WHITE)
    c.fill = PatternFill('solid', fgColor=INDIGO)
    c.alignment = Alignment(horizontal='center',vertical='center',wrap_text=True)
    c.border = box
r.row_dimensions[HR].height = 32

GOC = f'IF({IN}!$B$19<=0,0,{IN}!$B$15/{IN}!$B$19)'
for n in range(1,11):
    i = HR + n
    r.cell(row=i,column=1,value=n)
    r.cell(row=i,column=2,value=f'=ROUND({IN}!$B$22*MAX(0,12-{IN}!$B$23)*(1+{IN}!$B$28)^(A{i}-1),0)')
    r.cell(row=i,column=3,value=f'=ROUND({IN}!$B$24*12*(1+{IN}!$B$29)^(A{i}-1),0)')
    r.cell(row=i,column=4,value=f'=ROUND(B{i}*{IN}!$B$25,0)')
    r.cell(row=i,column=5,value=f'=B{i}-C{i}-D{i}')
    r.cell(row=i,column=6,value=f'=MAX(0,{IN}!$B$15-{GOC}*(A{i}-1))')
    r.cell(row=i,column=7,value=f'=MIN({GOC},F{i})')
    r.cell(row=i,column=8,value=f'=ROUND((F{i}+(F{i}-G{i}))/2*IF(A{i}=1,{IN}!$B$17,{IN}!$B$18),0)')
    r.cell(row=i,column=9,value=f'=G{i}+H{i}')
    r.cell(row=i,column=10,value=f'=E{i}-I{i}')
    r.cell(row=i,column=11,value=f'=IF(I{i}<=0,"—",ROUND(E{i}/I{i},2))')
    r.cell(row=i,column=12,value=f'=J{i}' if n==1 else f'=L{i-1}+J{i}')
    r.cell(row=i,column=13,value=f'=IF({IN}!$B$7<=0,"—",E{i}/{IN}!$B$7)')
    r.cell(row=i,column=14,value=f'=IF({IN}!$B$11<=0,"—",E{i}/{IN}!$B$11)')
    for col in range(1,15):
        c = r.cell(row=i,column=col)
        c.border = box
        c.font = Font(name='Montserrat',size=10,color=TDARK)
        c.alignment = Alignment(horizontal='right' if col>1 else 'center',vertical='center')
        if col in (2,3,4,5,6,7,8,9,10,12): c.number_format = M0
        if col == 11: c.number_format = NUM2
        if col in (13,14): c.number_format = PCT
        if col == 10: c.font = Font(name='Montserrat',size=10,bold=True,color=INDIGO)
        if n % 2 == 0: c.fill = PatternFill('solid', fgColor=CREAM)
    r.row_dimensions[i].height = 19

TR = HR + 11
r.cell(row=TR,column=1,value="TỔNG")
for col in [2,3,4,5,7,8,9,10]:
    L = get_column_letter(col)
    r.cell(row=TR,column=col,value=f'=SUM({L}{HR+1}:{L}{HR+10})')
for col in range(1,15):
    c = r.cell(row=TR,column=col)
    c.fill = PatternFill('solid', fgColor=GOLDL)
    c.font = Font(name='Montserrat',size=10,bold=True,color=TDARK)
    c.border = box
    c.alignment = Alignment(horizontal='right' if col>1 else 'center',vertical='center')
    if col in (2,3,4,5,7,8,9,10): c.number_format = M0
r.row_dimensions[TR].height = 22

r.merge_cells(start_row=TR+2,start_column=1,end_row=TR+3,end_column=14)
nt = r.cell(row=TR+2,column=1,value=("Cách tính trả nợ: gốc đều mỗi năm, lãi tính trên dư nợ bình quân trong năm — đúng cách các ngân hàng Việt Nam đang áp dụng cho vay mua bất động sản.  "
    "Năm 1 dùng lãi suất ưu đãi, từ năm 2 dùng lãi suất thả nổi.  ·  DAPANO GROUP · Think Kind!"))
nt.font = Font(name='Montserrat',size=8,italic=True,color=TMUTE)
nt.alignment = Alignment(wrap_text=True,vertical='top',indent=1)
r.freeze_panes = "A15"
r.sheet_view.showGridLines = False
r.protection.sheet = True; r.protection.enable()

# ═══════════════════ SHEET 3 · BA KỊCH BẢN ═══════════════════
s = wb.create_sheet("3. Ba kịch bản")
title_block(s,"BA KỊCH BẢN","Cùng một toà nhà, ba cách đời có thể diễn ra. Ô vàng sửa được.",11)
for i,w_ in enumerate([26,15,15,14,13,14,12,12,13,15,9],1):
    s.column_dimensions[get_column_letter(i)].width = w_

section(s,4,"GIẢ ĐỊNH LỆCH NHAU Ở ĐÂU",11)
s.cell(row=5,column=1,value="Điều chỉnh so với sheet Nhập liệu")
for j,name in enumerate(["THẬN TRỌNG","CƠ SỞ","KỲ VỌNG"]):
    c = s.cell(row=5,column=2+j,value=name)
    c.font = Font(name='Montserrat',size=10,bold=True,color=WHITE)
    c.fill = PatternFill('solid', fgColor=[MAGENTA,INDIGO,PURPLE][j])
    c.alignment = Alignment(horizontal='center',vertical='center')
s.cell(row=5,column=1).font = Font(name='Montserrat',size=10,bold=True,color=TDARK)
adj = [
 (6,"Giá thuê lệch",[-0.10,0,0.07],PCT),
 (7,"Số tháng trống mỗi năm",[f'={IN}!$B$23+1.5',f'={IN}!$B$23',f'=MAX(0,{IN}!$B$23-0.5)'],'0.0'),
 (8,"Lãi suất cộng thêm",[0.02,0,-0.005],PCT),
 (9,"Tăng giá thuê mỗi năm",[f'=MAX(0,{IN}!$B$28-0.01)',f'={IN}!$B$28',f'={IN}!$B$28+0.01'],PCT),
]
for row_i,lbl,vals,fmt in adj:
    c = s.cell(row=row_i,column=1,value=lbl)
    c.font = Font(name='Montserrat',size=10,color=TDARK); c.alignment = Alignment(indent=1,vertical='center')
    for j,v in enumerate(vals):
        cc = s.cell(row=row_i,column=2+j,value=v)
        cc.number_format = fmt; cc.border = gold_box
        cc.fill = PatternFill('solid', fgColor=GOLDL)
        cc.font = Font(name='Montserrat',size=10,bold=True,color=TDARK)
        cc.alignment = Alignment(horizontal='center',vertical='center')
        cc.protection = Protection(locked=False)

BHDR = ["Năm","Doanh thu thuê","Chi phí vận hành","Thuế cho thuê","NOI","Dư nợ đầu năm",
        "Trả gốc","Trả lãi","Tổng trả nợ","Dòng tiền ròng","DSCR"]
blocks = [(12,"KỊCH BẢN 1 · THẬN TRỌNG — thuê thấp, trống nhiều, lãi suất tăng","B",MAGENTA),
          (26,"KỊCH BẢN 2 · CƠ SỞ — đúng như số đã nhập","C",INDIGO),
          (40,"KỊCH BẢN 3 · KỲ VỌNG — thuê tốt, ít trống, lãi suất dễ thở","D",PURPLE)]
for start,ttl,X,color in blocks:
    section(s,start,ttl,11)
    for col in range(1,12): s.cell(row=start,column=col).fill = PatternFill('solid', fgColor=color)
    hr = start+1
    for i,h in enumerate(BHDR,1):
        c = s.cell(row=hr,column=i,value=h)
        c.font = Font(name='Montserrat',size=9,bold=True,color=WHITE)
        c.fill = PatternFill('solid', fgColor=INDIGO); c.border = box
        c.alignment = Alignment(horizontal='center',vertical='center',wrap_text=True)
    s.row_dimensions[hr].height = 30
    for n in range(1,11):
        i = hr+n
        s.cell(row=i,column=1,value=n)
        s.cell(row=i,column=2,value=f'=ROUND({IN}!$B$22*(1+${X}$6)*MAX(0,12-${X}$7)*(1+${X}$9)^(A{i}-1),0)')
        s.cell(row=i,column=3,value=f'=ROUND({IN}!$B$24*12*(1+{IN}!$B$29)^(A{i}-1),0)')
        s.cell(row=i,column=4,value=f'=ROUND(B{i}*{IN}!$B$25,0)')
        s.cell(row=i,column=5,value=f'=B{i}-C{i}-D{i}')
        s.cell(row=i,column=6,value=f'=MAX(0,{IN}!$B$15-{GOC}*(A{i}-1))')
        s.cell(row=i,column=7,value=f'=MIN({GOC},F{i})')
        s.cell(row=i,column=8,value=f'=ROUND((F{i}+(F{i}-G{i}))/2*(IF(A{i}=1,{IN}!$B$17,{IN}!$B$18)+${X}$8),0)')
        s.cell(row=i,column=9,value=f'=G{i}+H{i}')
        s.cell(row=i,column=10,value=f'=E{i}-I{i}')
        s.cell(row=i,column=11,value=f'=IF(I{i}<=0,"—",ROUND(E{i}/I{i},2))')
        for col in range(1,12):
            c = s.cell(row=i,column=col); c.border = box
            c.font = Font(name='Montserrat',size=10,color=TDARK)
            c.alignment = Alignment(horizontal='right' if col>1 else 'center',vertical='center')
            if col in range(2,11): c.number_format = M0
            if col == 11: c.number_format = NUM2
            if col == 10: c.font = Font(name='Montserrat',size=10,bold=True,color=INDIGO)
            if n % 2 == 0: c.fill = PatternFill('solid', fgColor=CREAM)
        s.row_dimensions[i].height = 18
    tr = hr+11
    s.cell(row=tr,column=1,value="TỔNG")
    for col in [2,3,4,5,7,8,9,10]:
        L = get_column_letter(col)
        s.cell(row=tr,column=col,value=f'=SUM({L}{hr+1}:{L}{hr+10})')
    for col in range(1,12):
        c = s.cell(row=tr,column=col)
        c.fill = PatternFill('solid', fgColor=GOLDL); c.border = box
        c.font = Font(name='Montserrat',size=10,bold=True,color=TDARK)
        c.alignment = Alignment(horizontal='right' if col>1 else 'center',vertical='center')
        if col in (2,3,4,5,7,8,9,10): c.number_format = M0

section(s,55,"SO SÁNH BA KỊCH BẢN",11)
cmp_rows = [
 ("Dòng tiền ròng năm 1 (triệu)",["=J14","=J28","=J42"],M0),
 ("Dòng tiền bình quân tháng năm 1",["=J14/12","=J28/12","=J42/12"],M0),
 ("DSCR năm 1",["=IF(I14<=0,\"—\",E14/I14)","=IF(I28<=0,\"—\",E28/I28)","=IF(I42<=0,\"—\",E42/I42)"],NUM2),
 ("Tổng dòng tiền 10 năm (triệu)",["=J24","=J38","=J52"],M0),
 ("Năm dòng tiền chuyển dương",['=IFERROR(MATCH(TRUE,INDEX($J$14:$J$23>0,0),0),"—")',
                                 '=IFERROR(MATCH(TRUE,INDEX($J$28:$J$37>0,0),0),"—")',
                                 '=IFERROR(MATCH(TRUE,INDEX($J$42:$J$51>0,0),0),"—")'],'0'),
]
s.cell(row=56,column=1,value="Chỉ số")
for j,name in enumerate(["THẬN TRỌNG","CƠ SỞ","KỲ VỌNG"]):
    c = s.cell(row=56,column=2+j,value=name)
    c.font = Font(name='Montserrat',size=10,bold=True,color=WHITE)
    c.fill = PatternFill('solid', fgColor=[MAGENTA,INDIGO,PURPLE][j])
    c.alignment = Alignment(horizontal='center',vertical='center')
s.cell(row=56,column=1).font = Font(name='Montserrat',size=10,bold=True,color=TDARK)
for k,(lbl,fs,fmt) in enumerate(cmp_rows):
    i = 57+k
    c = s.cell(row=i,column=1,value=lbl)
    c.font = Font(name='Montserrat',size=10,color=TDARK); c.alignment = Alignment(indent=1,vertical='center')
    for j,f in enumerate(fs):
        cc = s.cell(row=i,column=2+j,value=f)
        cc.number_format = fmt; cc.border = box
        cc.fill = PatternFill('solid', fgColor=BGSOFT)
        cc.font = Font(name='Montserrat',size=10,bold=True,color=MAGENTA)
        cc.alignment = Alignment(horizontal='center',vertical='center')
    s.row_dimensions[i].height = 19

s.merge_cells(start_row=63,start_column=1,end_row=64,end_column=11)
nn = s.cell(row=63,column=1,value=("Nếu kịch bản THẬN TRỌNG vẫn cho dòng tiền dương và DSCR trên 1,2 thì toà nhà này chịu được sóng gió. "
    "Nếu chỉ kịch bản KỲ VỌNG mới đẹp, nghĩa là anh/chị đang mua một hy vọng chứ không mua một dòng tiền.  ·  DAPANO GROUP"))
nn.font = Font(name='Montserrat',size=9,italic=True,color=TSOFT)
nn.fill = PatternFill('solid', fgColor=BGSOFT)
nn.alignment = Alignment(wrap_text=True,vertical='top',indent=1)
s.sheet_view.showGridLines = False
s.protection.sheet = True; s.protection.enable()

# ═══════════════════ SHEET 4 · ĐỌC KẾT QUẢ ═══════════════════
d = wb.create_sheet("4. Đọc kết quả")
title_block(d,"ĐỌC KẾT QUẢ BẰNG TIẾNG NGƯỜI","Năm chỉ số quyết định nên mua hay không",3)
d.column_dimensions['A'].width = 22
d.column_dimensions['B'].width = 62
d.column_dimensions['C'].width = 34
items = [
 ("NOI","Lãi ròng từ việc cho thuê trong một năm, sau khi trừ chi phí vận hành và thuế, nhưng CHƯA trừ tiền trả ngân hàng. Đây là sức khoẻ thật của toà nhà, không phụ thuộc anh/chị vay nhiều hay ít.","NOI âm nghĩa là toà nhà tự nó đã lỗ. Không có cách vay nào cứu được."),
 ("Cap Rate","NOI chia cho số tiền bỏ ra. Trả lời câu: mỗi đồng vốn mang về bao nhiêu một năm. Dùng để so hai toà nhà khác giá với nhau.","So Cap Rate với lãi gửi tiết kiệm. Thấp hơn tiết kiệm mà vẫn mua thì phải có lý do khác."),
 ("DSCR","NOI chia cho tổng tiền trả ngân hàng trong năm. Trả lời câu: tiền thuê có gánh nổi nợ không.","Dưới 1,2 là ĐÈN ĐỎ — chỉ cần trống hai tháng là phải bù tiền túi.\n1,2–1,5 là đèn vàng.\nTrên 1,5 là thở được."),
 ("Dòng tiền ròng","Số tiền thật sự còn lại trong túi sau khi đã trả ngân hàng. Đây là con số anh/chị tiêu được.","Âm trong 12–18 tháng đầu là bình thường nếu có cải tạo. Âm quá 18 tháng là phải xem lại."),
 ("Điểm hoà vốn","Năm mà dòng tiền luỹ kế bù đủ phần vốn tự có đã bỏ ra.","Toà nhà dòng tiền tốt thường thu hồi vốn tự có trong 7–10 năm, chưa tính phần tăng giá tài sản."),
]
row = 4
hd = ["Chỉ số","Nghĩa là gì","Ngưỡng cảnh báo"]
for i,h in enumerate(hd,1):
    c = d.cell(row=row,column=i,value=h)
    c.font = Font(name='Montserrat',size=10,bold=True,color=WHITE)
    c.fill = PatternFill('solid', fgColor=INDIGO); c.border = box
    c.alignment = Alignment(horizontal='center',vertical='center')
row = 5
for name,mean,warn in items:
    c1 = d.cell(row=row,column=1,value=name)
    c1.font = Font(name='Montserrat',size=11,bold=True,color=MAGENTA)
    c1.alignment = Alignment(vertical='center',horizontal='center')
    c2 = d.cell(row=row,column=2,value=mean)
    c2.font = Font(name='Montserrat',size=10,color=TDARK)
    c2.alignment = Alignment(wrap_text=True,vertical='top',indent=1)
    c3 = d.cell(row=row,column=3,value=warn)
    c3.font = Font(name='Montserrat',size=9,color=TSOFT)
    c3.alignment = Alignment(wrap_text=True,vertical='top',indent=1)
    for col in range(1,4):
        d.cell(row=row,column=col).border = box
        if row % 2 == 1: d.cell(row=row,column=col).fill = PatternFill('solid', fgColor=CREAM)
    d.row_dimensions[row].height = 78
    row += 1

row += 1
d.merge_cells(start_row=row,start_column=1,end_row=row,end_column=3)
c = d.cell(row=row,column=1,value="BA CÂU HỎI TỰ ĐẶT TRƯỚC KHI ĐẶT CỌC")
c.font = Font(name='Montserrat',size=11,bold=True,color=WHITE)
c.fill = PatternFill('solid', fgColor=PURPLE)
c.alignment = Alignment(indent=1,vertical='center'); d.row_dimensions[row].height = 24
row += 1
qs = ["1.  Nếu toà nhà để trống ba tháng liền, em có bù được không?",
      "2.  Nếu lãi suất thả nổi tăng thêm 2%, DSCR còn trên 1,2 không? (xem kịch bản Thận trọng)",
      "3.  Nếu năm năm nữa mới bán được, dòng tiền trong năm năm đó có đủ nuôi khoản vay không?"]
for q in qs:
    d.merge_cells(start_row=row,start_column=1,end_row=row,end_column=3)
    c = d.cell(row=row,column=1,value=q)
    c.font = Font(name='Montserrat',size=10,color=TDARK)
    c.alignment = Alignment(indent=2,vertical='center'); d.row_dimensions[row].height = 22
    row += 1

row += 1
d.merge_cells(start_row=row,start_column=1,end_row=row+3,end_column=3)
c = d.cell(row=row,column=1,value=("Ba con số đẹp trên file này chưa phải là một quyết định đúng. Còn pháp lý, còn hiện trạng công trình, còn khách thuê.\n\n"
    "Nếu anh/chị muốn em cùng ngồi xuống đọc lại từng dòng cho chính toà nhà anh/chị đang nhắm — em luôn sẵn sàng, và không mất phí.\n\n"
    "Cộng Hưởng Cùng Viên Mãn  ·  Nghĩ Thiện!  ·  DAPANO GROUP"))
c.font = Font(name='Montserrat',size=10,color=TSOFT)
c.fill = PatternFill('solid', fgColor=BGSOFT)
c.alignment = Alignment(wrap_text=True,vertical='top',indent=1)
d.sheet_view.showGridLines = False
d.protection.sheet = True; d.protection.enable()

wb.save("ke-hoach-tang-qua/tang-1/may-tinh-dong-tien.xlsx")
print("saved")
